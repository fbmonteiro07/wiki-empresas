from pathlib import Path
import json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[3]
OUT = ROOT / '_deliverables' / 'CPU_market_share_2021_2030_2026-09-23.docx'
OUT.parent.mkdir(exist_ok=True)
D = json.loads((BASE / 'research.json').read_text(encoding='utf-8'))
pc = D['pc_value']
da = (pc['2028'][1] - pc['2025'][1]) / 3
dr = (pc['2028'][2] - pc['2025'][2]) / 3
for year in [2029, 2030]:
    a = pc['2028'][1] + (year - 2028) * da
    r = pc['2028'][2] + (year - 2028) * dr
    pc[str(year)] = [100-a-r, a, r]
unit30 = [x / sum(D['server_unit_inputs_2030_mn']) * 100 for x in D['server_unit_inputs_2030_mn']]
checks = {}
for name in ['server_value', 'pc_value', 'server_units', 'pc_units', 'latest_quarter_units']:
    checks[name] = {'bounds_pass': all(0 <= x <= 100 for row in D[name].values() for x in row),
                    'sum_pass': all(abs(sum(row)-100) <= .11 for row in D[name].values()),
                    'max_sum_error_pp': max(abs(sum(row)-100) for row in D[name].values())}
assert all(c['bounds_pass'] and c['sum_pass'] for c in checks.values())
(BASE / 'calculations.json').write_text(json.dumps({'pc_extended':pc, 'server_unit_2030':unit30, 'pc_annual_steps_pp':[da,dr], 'checks':checks}, indent=2), encoding='utf-8')

# Print-quality chart drawn from the same data used in the editable Word tables.
font_path = 'C:/Windows/Fonts/segoeui.ttf'
bold_path = 'C:/Windows/Fonts/segoeuib.ttf'
def font(n, bold=False): return ImageFont.truetype(bold_path if bold else font_path,n)
W,H=2100,1380
im=Image.new('RGB',(W,H),'white'); g=ImageDraw.Draw(im)
COL=['#1F4E79','#B64926','#23786C']
def line(points, color, width=7, dash=None):
    if not dash: g.line(points, fill=color, width=width); return
    for p1,p2 in zip(points,points[1:]):
        dx,dy=p2[0]-p1[0],p2[1]-p1[1]; length=math.hypot(dx,dy)
        for a in range(0,int(length),dash*2):
            b=min(a+dash,length)
            g.line([(p1[0]+dx*a/length,p1[1]+dy*a/length),(p1[0]+dx*b/length,p1[1]+dy*b/length)],fill=color,width=width)
for k,(title,data) in enumerate([('Server CPU value share',D['server_value']),('PC CPU value share',pc)]):
    top=50+k*650; left=130; right=1775; bottom=top+480
    g.text((left,top-18),title,fill='black',font=font(42,True))
    ytop=top+85
    x=lambda yr: left+(yr-2021)*(right-left)/9
    y=lambda val: bottom-(bottom-ytop)*val/100
    g.rectangle((x(2025.5),ytop,right,bottom),fill='#F0F3F5')
    if k==1: g.rectangle((x(2028.5),ytop,right,bottom),fill='#FFF3DC')
    for v in [0,25,50,75,100]:
        yy=y(v); g.line([(left,yy),(right,yy)],fill='#D9DFE4',width=2)
        g.text((left-95,yy-20),f'{v}%',fill='#5E666D',font=font(29))
    for yr in range(2021,2031):
        g.text((x(yr)-36,bottom+18),str(yr)[2:]+('E' if yr>=2026 else ''),fill='#3D444B',font=font(28))
    for j,color in enumerate(COL):
        pts=[(x(yr),y(data[str(yr)][j])) for yr in range(2021,2031)]
        line(pts[:5],color)
        line(pts[4:8] if k==1 else pts[4:],color,dash=19)
        if k==1: line(pts[7:],color,width=7,dash=5)
        for px,py in pts: g.ellipse((px-6,py-6,px+6,py+6),fill=color)
        g.text((right+28,y(data['2030'][j])-20),f"{data['2030'][j]:.1f}%",font=font(32,True),fill=color)
    g.text((x(2026)-35,ytop-45),'Broker forecast',fill='#636B72',font=font(27))
    if k==1: g.text((x(2029)-48,ytop-45),'Extension',fill='#956720',font=font(27))
    for j,label in enumerate(['Intel','AMD','Arm-based']):
        xx=left+j*320
        g.line([(xx,bottom+94),(xx+55,bottom+94)],fill=COL[j],width=8)
        g.text((xx+72,bottom+73),label,fill='#30373D',font=font(31))
im.save(BASE/'share_chart.png')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.64); sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
sec.footer_distance=Inches(.28)
styles=doc.styles
for nm in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=styles[nm]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.08
styles['Title'].font.size=Pt(28); styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(5)
styles['Subtitle'].font.size=Pt(12)
styles['Subtitle'].font.italic=False
styles['Subtitle'].paragraph_format.space_after=Pt(15)
for style in styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
styles['Heading 1'].font.size=Pt(19); styles['Heading 1'].font.bold=True
styles['Heading 1'].paragraph_format.space_before=Pt(0)
styles['Heading 1'].paragraph_format.space_after=Pt(10)
styles['Heading 2'].font.size=Pt(13); styles['Heading 2'].font.bold=True
styles['Heading 2'].paragraph_format.space_before=Pt(9)
styles['Heading 2'].paragraph_format.space_after=Pt(5)
for nm in ['Caption']:
    styles[nm].font.name='Calibri'; styles[nm].font.size=Pt(9)
    styles[nm].font.color.rgb=RGBColor.from_string('4B5560')
    styles[nm].paragraph_format.space_after=Pt(7)

def para(text, bold=False, style=None):
    p=doc.add_paragraph(style=style)
    r=p.add_run(text); r.bold=bold
    return p
def heading(text, level=1): doc.add_heading(text,level)
def note(text): return para(text,style='Caption')
def page(): doc.add_page_break()
def shade(cell,fill):
    e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); cell._tc.get_or_add_tcPr().append(e)
def table(headers, rows, widths, small=10.5, status_col=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    pr=t._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+tag); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
    pr.append(b)
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for ci,(c,txt) in enumerate(zip(t.rows[0].cells,headers)):
        c.text=txt; c.width=Inches(widths[ci]); shade(c,'243E50')
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci,(c,value) in enumerate(zip(cells,row)):
            c.text=str(value); c.width=Inches(widths[ci])
            if ri%2==0:shade(c,'F2F5F7')
            if status_col is not None and ci==status_col and str(value)=='X':shade(c,'FFF0CE')
        keep=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(keep)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for tag,val in [('top','62'),('bottom','62'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+tag);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');m.append(e)
            tcPr.append(m)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or widths[ci]>3 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(small)
                    if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def link(p,text,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245C85');pr.append(c);r.append(pr)
    tt=OxmlElement('w:t');tt.text=text;r.append(tt);h.append(r);p._p.append(h)
def archive_link(code,title,path,pages):
    p=para(code+' ');p.runs[0].bold=True
    link(p,title,Path(path).as_uri());p.add_run(' '+pages)

foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('CPU market share  |  ');r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('68727A')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
doc.core_properties.title='CPU market share 2021 to 2030'
doc.core_properties.subject='PC and server CPU historical value shares and projections with unit-share comparisons'
doc.core_properties.author='Capstone research support'

# PAGE 1
para('CPU market share 2021 to 2030',style='Title')
para('PC and server CPUs   |   Five historical years and five projected years',style='Subtitle')
para('BofA projects Arm-based processors to lead server CPU value share by 2030 at 47.3%, versus AMD at 30.7% and Intel at 22.0%. Intel remains the largest PC CPU supplier in the illustrative 2030 extension, at 52.7%. [S1, pp. 6 and 21; X1]')
para('Scope: worldwide PC and server CPU silicon, shown separately. The main measure is annual value share, including x86 and Arm-based designs. History covers 2021–2025; projections cover 2026–2030, including the current year. Prepared 23 September 2026. Unit share is shown on page 3.')
doc.add_picture(str(BASE/'share_chart.png'),width=Inches(7.02))
note('Source: BofA / Vivek Arya et al., 12 Aug 2026, pp. 6 and 21. Solid lines are historical estimates; dashed lines are broker forecasts. Dotted PC lines for 2029–2030 are the illustrative extension X1, calculated 23 Sep 2026. All figures are percentages.')
para('Read value share with care',bold=True)
para('Arm-based covers multiple suppliers and custom chips; it is not Arm Holdings revenue. Value shares include imputed custom-silicon values. BofA revised Arm’s 2025 server value share from 11.0% in February to 32.3% in August, while its 2025 unit share changed only from 13.6% to 13.7%. This is a restatement for the same year, not adoption between report dates. [S4, p. 17; S1, pp. 5 and 19]')

# PAGE 2
page();heading('Annual CPU value share')
para('Each row measures the same annual CPU market within its segment. H = historical estimate; F = BofA forecast; X = illustrative extension. Values may total 99.9% or 100.1% because of source rounding.')
heading('Server CPUs',2)
rows=[]
for yr in range(2021,2031):rows.append([str(yr),*[f'{x:.1f}%' for x in D['server_value'][str(yr)]],'H' if yr<=2025 else 'F'])
table(['Calendar year','Intel','AMD','Arm-based','Status'],rows,[1.30,1.38,1.38,1.60,1.38],status_col=4)
note('S1: BofA / Arya et al., 12 Aug 2026, p. 21 for 2021–2028 and p. 6 for 2029–2030. Arm combines merchant and custom processors; value includes estimates for captive silicon.')
heading('PC CPUs',2)
rows=[]
for yr in range(2021,2031):rows.append([str(yr),*[f'{x:.1f}%' for x in pc[str(yr)]],'H' if yr<=2025 else ('F' if yr<=2028 else 'X')])
table(['Calendar year','Intel','AMD','Arm-based','Status'],rows,[1.30,1.38,1.38,1.60,1.38],status_col=4)
note('S1: BofA / Arya et al., 12 Aug 2026, p. 21 through 2028. X1: 2029–2030 are illustrative extensions of the 2025–2028 average annual share changes, without independent empirical validation. They are not BofA forecasts; see page 4.')

# PAGE 3
page();heading('Unit share and forecast comparisons')
para('Shipment share measures CPU units, rather than the assigned value of those units. It is also different from installed-base share, CPU core share and complete-server-system share. An Arm architecture share should not be compared directly with an AMD share restricted to x86.')
heading('Annual unit share in the BofA model',2)
rows=[]
for segment,key in [('Server','server_units'),('PC','pc_units')]:
    for yr in range(2023,2027): rows.append([segment,str(yr)+('E' if yr==2026 else ''),*[f'{v:.1f}%' for v in D[key][str(yr)]]])
table(['Segment','Calendar year','Intel','AMD','Arm-based'],rows,[1.15,1.40,1.38,1.38,1.73])
note('S1: BofA / Arya et al., 12 Aug 2026, pp. 16 and 19; Mercury Research, company reports and BofA estimates. 2023–2025 are historical estimates; 2026E is a full-year forecast. The value table uses p. 21; see the PC source discrepancy on page 5.')
heading('Latest quarterly checkpoint',2)
table(['Second quarter 2026','Intel','AMD','Arm-based'],[
    ['Server units', '53.9%','28.4%','17.6%'],['PC units','59.1%','25.6%','15.3%']], [2.32,1.45,1.45,1.82])
note('S1: Mercury Research as reproduced by BofA / Arya et al., 12 Aug 2026, pp. 16 and 19. Quarterly readings are separate from annual estimates.')
heading('Server CPU unit share in 2030',2)
table(['Forecast vintage','Intel','AMD','Arm-based'],[
    ['BofA 12 Aug 2026',*[f'{v:.1f}%' for v in unit30]],
    ['UBS 21 May 2026','29.0%','29.0%','42.0%']], [2.32,1.45,1.45,1.82])
para('BofA’s unit shares above are calculated from 29.3 million Intel, 24.9 million AMD and 27.8 million Arm-based CPUs, totaling 82.0 million. For Arm: 27.8 ÷ 82.0 × 100 = 33.9%. UBS publishes the shares directly. [S1, p. 5; S2, p. 2]')
para('This is a comparison of two research models, not consensus. Their starting market sizes differ: BofA has 29.9 million server CPUs in 2025; UBS has 23 million. Their 2030 Arm unit-share difference should therefore be treated as a model and scope debate. [S1, p. 5; S2, p. 2]')

# PAGE 4
page();heading('PC extension and sensitivity')
para('X1 extends BofA’s PC value-share table by two years. It carries forward the average annual percentage-point gains for AMD and Arm from 2025 to 2028, then calculates Intel as the remainder. This is a mechanical scenario, not a separate demand, pricing or product forecast.')
heading('Worked calculation for 2030',2)
para('AMD annual gain = (25.8% − 23.2%) ÷ 3 = 0.8667 percentage points. Arm annual gain = (16.8% − 12.3%) ÷ 3 = 1.5000 percentage points. The inputs are BofA’s 2025 historical estimates and 2028 forecasts. [S1, p. 21]')
para('AMD 2030 = 25.8% + 2 × 0.8667 = 27.5%. Arm 2030 = 16.8% + 2 × 1.5000 = 19.8%. Intel 2030 = 100% − 27.5333% − 19.8% = 52.7%. Calculations retain precision until display. [X1, 23 Sep 2026]')
heading('2030 PC value share under different assumptions',2)
table(['Post 2028 share gains','Intel','AMD','Arm-based'],[
    ['0× annual gains','57.4%','25.8%','16.8%'],
    ['1× annual gains used','52.7%','27.5%','19.8%'],
    ['1.5× annual gains','50.3%','28.4%','21.3%']], [2.32,1.45,1.45,1.82])
note('X1, calculated 23 Sep 2026 from S1 p. 21. The multipliers are illustrative assumptions, not probabilities or confidence bounds. Holding the other series constant, a 0.5-point change in either annual gain changes its 2030 share by 1 point and Intel’s share by the opposite amount.')
heading('What is sourced and what is assumed',2)
t=table(['Input tag','Meaning in this document'],[
    ['HARD','Exact source values: the historical estimates and published broker forecasts. Sourced does not mean certain.'],
    ['PARTIAL','Derived quantities: annual share-gain rates and BofA 2030 unit shares, calculated from published rounded inputs.'],
    ['ESTIMATE','The assumption that the PC share-gain rates persist after 2028 and the sensitivity multipliers.']], [1.25,5.79],small=10.5)
for i,fill in [(1,'E4F0E6'),(2,'E5EFF9'),(3,'FFF0CE')]: shade(t.rows[i].cells[0],fill)
heading('Validation limits',2)
para('Arithmetic checks pass: every share is between 0% and 100%; all displayed source rows sum to 100% within 0.1 point; the extension sums to 100% before rounding. All market-share comparisons use annual flows or explicitly labeled quarterly flows.')
para('The PC extension is anchored to a historical estimate and a future broker estimate. It has no independent empirical validation or backtest. The 2030 figures are directional. Independent source audit: PASS on transcription and arithmetic; PARTIAL on the model because future outcomes and the PC extension are not externally validated.')

# PAGE 5
page();heading('Sources and interpretation')
heading('Primary research used',2)
archive_link('S1', 'BofA Global Research — Rise of the agents raising CPU TAM again to 210bn.',D['sources']['S1']['path'], 'Vivek Arya, Duksan Jang, Michael Mani and Liam Pharr; 12 August 2026.')
para('Page 5, Exhibit 3: server vendor shipments, value shares and pricing assumptions. Page 6, Exhibit 4: server value-share forecast through 2030. Pages 16 and 19, Exhibits 17 and 26: PC and server shipment shares. Page 20, Exhibit 29: Mercury Research unit outlook. Page 21, Exhibit 30: annual PC and server value-share history and forecasts.')
archive_link('S2','UBS — Quantifying the server CPU opportunity.',D['sources']['S2']['path'],'Sunny Lin, Randy Abrams and Nicolas Gaudois; 21 May 2026; pp. 1–2, Figure 2. Used as a separately attributed server unit-share comparison.')
archive_link('S4','BofA Global Research — CPU in AI key inference beneficiary TAM could more than double by CY30.',str(ROOT/'relatórios bons'/'CPU.html'),'Vivek Arya et al.; 23 February 2026; p. 17. Used only to illustrate revisions to historical estimates, not as the forecast baseline.')
p=para('S3 ');p.runs[0].bold=True;link(p,'Mercury Research product methodology','https://www.mercuryresearch.com/products.shtml');p.add_run('. Accessed 23 September 2026. Describes shipment, pricing, revenue and segment-share coverage. No numerical market shares in this document are taken from that webpage.')
para('X1. Calculated on 23 September 2026 from S1 p. 21. The 2029–2030 PC extension and sensitivity cases are explicitly illustrative; they are not an approved Capstone house forecast.')
heading('Source differences that affect interpretation',2)
para('Historical value shares can change with the model. BofA’s February report estimated Arm’s 2025 server value share at 11.0%; its August report uses 32.3%. The corresponding unit shares are much closer, at 13.6% and 13.7%. The change between reports is a revision to the estimate for the same year, not a new shipment event. [S4, p. 17; S1, pp. 19 and 21]')
para('Within the August report, PC 2025 value shares differ between the annual summary on p. 21 (Intel 64.5%, AMD 23.2%, Arm 12.3%) and p. 16 (64.8%, 23.3%, 11.9%). The main tables and PC extension consistently use p. 21. The unit table separately reproduces the p. 16 shipment series. [S1]')
para('Custom silicon value is imputed. BofA states modeling assumptions of roughly 75% of AMD’s average price for custom Arm and 1.25 times for merchant Arm. These are not uniform realized ratios across the table: aggregate merchant Arm also reflects separate NVIDIA pricing and mix. Value share is therefore sensitive to those conventions. The Arm-based category includes licensees and custom designs, rather than just Arm Holdings. [S1, p. 5]')
para('The main paths retain one BofA forecast vintage. UBS is shown separately. Market-share changes alone do not establish revenue growth, returns or a stock recommendation; the report compares CPU silicon shares and excludes GPUs, smartphones and embedded processors.')

doc.save(OUT)
print(str(OUT))
print(json.dumps({'checks':checks, 'pc2030':pc['2030'], 'server2030units':unit30},indent=2))
