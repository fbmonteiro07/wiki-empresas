# -*- coding: utf-8 -*-
"""Slice the Mix / Regiões / Clientes sheet code out of build_ate_consolidado.py, apply the review fixes, and save as
ate_extra_sheets.py (exec'd by build_ate_fernanda.py with `wb` in scope)."""
import os, re
HERE = os.path.dirname(os.path.abspath(__file__))
src = open(os.path.join(HERE, "build_ate_consolidado.py"), encoding="utf-8").read()
start = src.index("# Sheet 2 — Mix e segmentos (history)")
end = src.index("# Sheet 5 — Fontes e notas")
body = src[start:end]
# helpers needed by the sliced code (copied verbatim from the old builder, with note_col)
helpers = '''
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
F_TITLE = Font(name="Arial", size=16, bold=True, color="FF203864"); F_TXT = Font(name="Arial", size=10, color="FF202020")
F_TXTB = Font(name="Arial", size=10, bold=True, color="FF202020"); F_NOTE = Font(name="Arial", size=9, color="FF666666")
F_IN = Font(name="Arial", size=10, color="FF0000FF"); F_LINK = Font(name="Arial", size=10, color="FF008000")
FILL_H = PatternFill("solid", fgColor="FFDCE6F1"); FILL_S = PatternFill("solid", fgColor="FF203864"); FILL_Y = PatternFill("solid", fgColor="FFFFF2CC")
NF_N = '#,##0;\\\\(#,##0\\\\);"-"'; NF_N1 = '#,##0.0;\\\\(#,##0.0\\\\);"-"'; NF_N2 = '0.00;\\\\(0.00\\\\);"-"'; NF_P = '0.0%;\\\\(0.0%\\\\);"-"'; NF_X = '0.0"x"'

def setc(ws, a, v, font=F_TXT, nf=None, fill=None, al=None, wrap=False):
    c = ws[a]; c.value = v; c.font = font
    if nf: c.number_format = nf
    if fill: c.fill = fill
    if al or wrap: c.alignment = Alignment(horizontal=al, wrap_text=wrap, vertical="top" if wrap else None)
    return c

def hdr(ws, r, text, c0="C", c1="K"):
    import openpyxl
    setc(ws, f"{c0}{r}", text, F_TXTB, fill=FILL_H)
    for ci in range(openpyxl.utils.column_index_from_string(c0) + 1, openpyxl.utils.column_index_from_string(c1) + 1):
        ws[f"{L(ci)}{r}"].fill = FILL_H

def yhdr(ws, r, labels, c0=4):
    for i, t in enumerate(labels):
        setc(ws, f"{L(c0 + i)}{r}", t, F_TXTB, fill=FILL_H, al="center")

def row(ws, r, label, vals, nf=NF_N, font=F_TXT, c0=4, note=None, bold=False, src_font=None, note_col=None):
    setc(ws, f"C{r}", label, F_TXTB if bold else F_TXT)
    for i, v in enumerate(vals):
        if v is None: continue
        f = font if not (isinstance(v, str) and v.startswith("=")) else (F_TXT if font is F_IN else font)
        if bold: f = Font(name="Arial", size=10, bold=True, color=f.color.rgb if f.color else "FF202020")
        setc(ws, f"{L(c0 + i)}{r}", v, f, nf=nf, al="right")
    if note: setc(ws, f"{note_col or L(c0 + len(vals) + 1)}{r}", note, F_NOTE, al="left")

def sheet(wb, name, title, sub, widths):
    ws = wb.create_sheet(name); ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2.25; ws.column_dimensions["B"].width = 2.25
    for k, w in widths.items(): ws.column_dimensions[k].width = w
    setc(ws, "C2", title, F_TITLE); setc(ws, "C3", sub, F_TXT); return ws
'''
fixes = [
 ('row(wm, r, "SoC: Computing / Communications (% do SoC)", [0.90, None, 0.95], NF_P, F_IN, note=', 'row(wm, r, "SoC: Computing / Communications (% do SoC)", [0.90, None, 0.95], NF_P, F_IN, note_col="R", note='),
 ('row(wm, r, "Memória: DRAM (% da memória)", [0.95, 0.90, 0.85], NF_P, F_IN, note=', 'row(wm, r, "Memória: DRAM (% da memória)", [0.95, 0.90, 0.85], NF_P, F_IN, note_col="R", note='),
 ('row(wm, r, "SoC: ~80% HPC / AI (comentário)", [None, 0.80, None], NF_P, F_IN, note=', 'row(wm, r, "SoC: ~80% HPC / AI (comentário)", [None, 0.80, None], NF_P, F_IN, note_col="R", note='),
 ('row(wm, r, "AI-related como % da receita", [0.50, 0.60, 0.70, None], NF_P, F_IN, note=', 'row(wm, r, "AI-related como % da receita", [0.50, 0.60, 0.70, None], NF_P, F_IN, note_col="R", note='),
 ('row(wm, r, "Compute como % da receita de produto SOC", [None, None, None, 0.70], NF_P, F_IN, note=', 'row(wm, r, "Compute como % da receita de produto SOC", [None, None, None, 0.70], NF_P, F_IN, note_col="R", note='),
 ('row(wm, r, "SOC / Memory / IST no trimestre (US$ mi)", [None, None, None, "843 / 212 / 67"], None, F_IN, note=', 'row(wm, r, "SOC / Memory / IST no trimestre (US$ mi)", [None, None, None, "843 / 212 / 67"], None, F_IN, note_col="R", note='),
 ('    row(wm, r, name, vals, NF_N1, F_IN, note=src); trow[name] = r; r += 1', '    row(wm, r, name, vals, NF_N1, F_IN, note=src, note_col="R"); trow[name] = r; r += 1'),
 ('(diferenças de arredondamento/corporate)."); r += 1', '(diferenças de arredondamento/corporate).", note_col="R"); r += 1'),
 ('fxr = r; row(wm, r, "FX ¥/US$ (média FY-Mar)", ADV_FX_FY, NF_N1, F_IN, note="BBG USDJPY médias mensais abr→mar; FY26e = premissa da companhia (¥152; YTD 159).", ); r += 1',
  'row(wm, r, "Receita total reportada (¥ bi) — referência", [282.5, 275.9, 312.8, 416.9, 560.2, 486.5, 779.7, 1128.6, 1714.0], NF_N1, F_IN, note="Tanshin / guidance jul-2026; a soma dos segmentos BBG difere por arredondamento (≤¥1,3 bi).", note_col="R"); r += 1\nfxr = r; row(wm, r, "FX ¥/US$ (média FY-Mar)", ADV_FX_FY, NF_N1, F_IN, note="BBG USDJPY médias mensais abr→mar; FY26e = premissa da companhia (¥152; YTD 159).", note_col="R"); r += 1'),
 ('note="2021-23 = ASML_Peers_SemiCap_Felipe.xlsx Industry Assumptions (base antiga, ~40% abaixo da série Advantest em 2025); 2024-25 deck Advantest 2026-07-29."); d6 = r; r += 1',
  'note="2021-23 = ASML_Peers_SemiCap_Felipe.xlsx Industry Assumptions (base antiga; o mesmo arquivo carrega 6,8 / 8,0 bi para 2025 / 2026, ~40% abaixo da série da Advantest); 2024-25 deck Advantest 2026-07-29.", note_col="R"); d6 = r; r += 1'),
 ('NF_P, F_TXT, note="2021-23 >100% confirma que a base antiga do peer model subestima o TAM."); r += 1',
  'NF_P, F_TXT, note="⚠ Numerador Advantest em FY-Mar (inclui jan-mar do ano seguinte) vs TAM CY — infla a razão em anos de crescimento (2024-25 ~96-99% é teto, não leitura). 2021-23 na base antiga: 69-84%.", note_col="R"); r += 1'),
 ('row(wr, r, "Japão + resto do mundo (residual)",', 'row(wr, r, "Resto do mundo (residual; 2025 inclui Japão ~2%)",'),
 ('note="2025: Japão ~2% + RoW ~4% (10-K %); o valor de Japão 2025 não foi extraído em US$.");', 'note="2025: Japão ~2% + RoW ~4% (10-K %); o valor de Japão 2025 não foi extraído em US$.", note_col="R");'),
 ("""NF_N1, F_IN, bold=True, note="10-K FY23 / FY24 / FY25 'revenues by country'; 10-Q 2Q FY26 (six months / three months ended 28-jun-2026)."); ttr = r""",
  """NF_N1, F_IN, bold=True, note="10-K FY23 / FY24 / FY25 'revenues by country'; 10-Q 2Q FY26 (six months / three months ended 28-jun-2026).", note_col="R"); ttr = r"""),
 ('note="Reportado: 138,7 / 190,5 / 218,2 / 232,3 / 263,8 / 262,9 / 273,8 / 328,1 / 367,5."); r += 1', 'note="Reportado: 138,7 / 190,5 / 218,2 / 232,3 / 263,8 / 262,9 / 273,8 / 328,1 / 367,5.", note_col="R"); r += 1'),
 ('row(wr, r, "  % Ásia", [f"=D{r-2}/D{r-1}", f"=E{r-2}/E{r-1}"], NF_P, F_TXT, note=', 'row(wr, r, "  % Ásia", [f"=D{r-2}/D{r-1}", f"=E{r-2}/E{r-1}"], NF_P, F_TXT, note_col="R", note='),
]
n = 0
for a, b in fixes:
    if a in body: body = body.replace(a, b); n += 1
    else: print("NOT FOUND:", a[:70])
out = "# generated by make_extra_sheets.py — Mix / Regiões / Clientes sheets (exec'd with `wb` in scope)\n" + helpers + "\nr = 5\n" + body
open(os.path.join(HERE, "ate_extra_sheets.py"), "w", encoding="utf-8").write(out)
print("ate_extra_sheets.py written;", n, "fixes applied of", len(fixes))
