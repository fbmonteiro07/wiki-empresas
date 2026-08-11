#!/usr/bin/env python3
"""Validate the selected-index 2024-2030 revenue/EBIT workbook."""

import argparse
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook


EXPECTED = ["README", "Inputs & Anchors", "Growth Summary", "By Index", "By Sector", "Company Model", "Guardrails", "Sources"]


def hmap(ws, row=4):
    return {ws.cell(row, c).value: c for c in range(1, ws.max_column + 1)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    audit = json.loads(args.audit.read_text(encoding="utf-8"))
    failures = []
    warnings = []
    checks = {}
    with zipfile.ZipFile(args.workbook) as archive:
        bad = archive.testzip()
        checks["xlsx_zip_integrity"] = bad is None
        if bad:
            failures.append(f"Corrupt XLSX member: {bad}")
    wb = load_workbook(args.workbook, data_only=False, read_only=False)
    checks["sheets"] = wb.sheetnames
    if wb.sheetnames != EXPECTED:
        failures.append(f"Unexpected sheet names/order: {wb.sheetnames}")
    summary = wb["Growth Summary"]
    sh = hmap(summary)
    for row, year in enumerate(range(2024, 2031), 5):
        if summary.cell(row, sh["Year"]).value != year:
            failures.append(f"Summary year mismatch at row {row}")
        expected_revenue = audit["summary"][str(year)]["revenue_usd_bn"]
        actual_revenue = summary.cell(row, sh["Revenue USD bn"]).value
        if abs(actual_revenue - expected_revenue) > 1e-6:
            failures.append(f"Revenue mismatch FY{year}")
        low_formula = summary.cell(row, sh["AI @ 2% USD bn"]).value
        high_formula = summary.cell(row, sh["AI @ 3% USD bn"]).value
        if "'Inputs & Anchors'!$B$5" not in str(low_formula) or "'Inputs & Anchors'!$B$6" not in str(high_formula):
            failures.append(f"AI formula reference mismatch FY{year}")
    company = wb["Company Model"]
    company_count = company.max_row - 4
    checks["company_rows"] = company_count
    if company_count != 4946:
        failures.append(f"Company rows {company_count} != 4946")
    ch = hmap(company)
    low_col = ch["AI_2030_2pct_USD_m"]
    high_col = ch["AI_2030_3pct_USD_m"]
    low_formulas = sum(isinstance(company.cell(r, low_col).value, str) and company.cell(r, low_col).value.startswith("=") for r in range(5, company.max_row + 1))
    high_formulas = sum(isinstance(company.cell(r, high_col).value, str) and company.cell(r, high_col).value.startswith("=") for r in range(5, company.max_row + 1))
    checks["company_ai_low_formulas"] = low_formulas
    checks["company_ai_high_formulas"] = high_formulas
    if low_formulas != 4946 or high_formulas != 4946:
        failures.append("Missing company-level AI formulas")
    formula_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("=") and "#REF!" in value:
                    formula_errors.append(f"{ws.title}!{cell.coordinate}")
    checks["ref_formula_errors"] = formula_errors
    if formula_errors:
        failures.append(f"#REF formulas: {formula_errors[:10]}")
    by_index = wb["By Index"]
    by_sector = wb["By Sector"]
    checks["by_index_rows"] = by_index.max_row - 4
    checks["by_sector_rows"] = by_sector.max_row - 4
    if by_index.max_row - 4 != 49:
        failures.append("By Index should have 7 universes x 7 years = 49 rows")
    sector_count = len(audit["terminal_growth"])
    if by_sector.max_row - 4 != sector_count * 7:
        failures.append("By Sector row count mismatch")
    s2030 = audit["summary"]["2030"]
    cagr = (s2030["revenue_usd_bn"] / audit["summary"]["2026"]["revenue_usd_bn"]) ** 0.25 - 1
    checks["revenue_2026_2030_cagr"] = cagr
    checks["ebit_margin_2030"] = s2030["weighted_ebit_margin"]
    checks["ai_2030_low_usd_bn"] = s2030["ai_low_usd_bn"]
    checks["ai_2030_high_usd_bn"] = s2030["ai_high_usd_bn"]
    if cagr > 0.10:
        warnings.append("2026-2030 CAGR above 10%; inspect IT terminal growth")
    if audit["terminal_growth"]["Information Technology"]["base"] > 0.20:
        warnings.append("Information Technology terminal growth exceeds 20% and is load-bearing")
    report = {"status": "PASS" if not failures else "FAIL", "failures": failures, "warnings": warnings, "checks": checks}
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
