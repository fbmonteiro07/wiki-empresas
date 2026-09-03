# -*- coding: utf-8 -*-
"""Word report for the sentiment backtest (hot/cold × bull/bear): weekly panels, lead/lag, weekly picks,
earnings event windows. Reads the JSON the dashboard reads; draws charts with PIL (no matplotlib here).

Usage: py build_sentiment_report.py [--out path.docx]      → default _wiki/_data/sentiment/Sentiment_Backtest_<date>.docx
"""
import sys, json, math, argparse, datetime as dt
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]; DATA = ROOT / "_data" / "sentiment"
TODAY = dt.date.today()
J = lambda n: json.load(open(DATA / n, encoding="utf-8"))
PW, PL, LL, PK, EV = J("panel_weekly.json"), J("panel_long.json"), J("leadlag.json"), J("picks.json"), J("events.json")
SS = J("sellside_daily.json") if (DATA / "sellside_daily.json").exists() else {}
TW = J("tweet_daily.json")

BLUE, RED, ORANGE, GRAY, INK, INK2, GRID = (42, 120, 214), (227, 73, 72), (235, 104, 52), (138, 136, 127), (11, 11, 11), (82, 81, 78), (232, 230, 225)
FONT = lambda sz, bold=False: ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf", sz)
S = 2  # supersample factor for crisp PNGs

def pct(x, d=1): return "–" if x is None else f"{x*100:+.{d}f}%"
def sgn(x, d=3): return "–" if x is None else f"{x:+.{d}f}"
def fmt(x, d=1): return "–" if x is None else f"{x:.{d}f}"

# ----------------------------------------------------------------------------- charts (PIL)
def chart_ccf(path):
    """Small multiples: IC by lag for every Bloomberg-track signal. k<0 = return before the signal, k>0 = after."""
    C = LL["long"]["ccf"]; V = LL["long"]["verdict"]; lab = LL["labels"]; lags = LL["lags"]; keys = list(C)
    cols = 2; rows = math.ceil(len(keys) / cols); cw, ch = 900 * S, 230 * S; pad = 30 * S
    img = Image.new("RGB", (cols * cw + pad, rows * ch + pad * 2), "white"); d = ImageDraw.Draw(img)
    d.text((pad, pad // 3), "Cross-correlation with weekly universe-relative returns, Jan-2024 → Aug-2026 (139 weeks). Bars left of the shaded k=0 column: sentiment reacting to price. Right: sentiment leading price. Faded: |t| < 2.", fill=INK2, font=FONT(15 * S))
    for n, k in enumerate(keys):
        cx, cy = pad + (n % cols) * cw, pad + 20 * S + (n // cols) * ch
        m = {"l": 70 * S, "r": 30 * S, "t": 34 * S, "b": 34 * S}; W, H = cw - 20 * S, ch - 10 * S
        c = C[k]; mx = max(0.06, *[abs((c.get(str(l)) or {}).get("ic", 0) or 0) for l in lags])
        bw = (W - m["l"] - m["r"]) / len(lags); Y = lambda v: cy + m["t"] + (mx - v) / (2 * mx) * (H - m["t"] - m["b"])
        d.text((cx + m["l"], cy + 6 * S), lab.get(k, k), fill=INK, font=FONT(15 * S, True))
        d.text((cx + W - m["r"], cy + 6 * S), V.get(k, ""), fill=INK2, font=FONT(13 * S), anchor="ra")
        i0 = lags.index(0); d.rectangle([cx + m["l"] + i0 * bw, cy + m["t"], cx + m["l"] + (i0 + 1) * bw, cy + H - m["b"]], fill=(243, 242, 238))
        d.line([cx + m["l"], Y(0), cx + W - m["r"], Y(0)], fill=GRAY, width=S)
        for i, l in enumerate(lags):
            o = c.get(str(l)) or {}; v = o.get("ic", 0) or 0; sig = abs(o.get("t") or 0) >= 2
            col = BLUE if v >= 0 else RED
            if not sig: col = tuple(int(0.45 * a + 0.55 * 255) for a in col)
            x0 = cx + m["l"] + i * bw + bw * 0.2; x1 = x0 + bw * 0.6
            d.rectangle([x0, min(Y(0), Y(v)), x1, max(Y(0), Y(v)) + (1 if v == 0 else 0)], fill=col)
            d.text(((x0 + x1) / 2, cy + H - m["b"] + 6 * S), f"{l:+d}" if l else "0", fill=INK2, font=FONT(12 * S), anchor="ma")
            d.text(((x0 + x1) / 2, (Y(v) - 4 * S) if v >= 0 else (Y(v) + 4 * S)), f"{v:+.3f}", fill=INK2, font=FONT(11 * S), anchor="mb" if v >= 0 else "ma")
        d.text((cx + m["l"] - 6 * S, Y(mx)), f"{mx:.2f}", fill=GRAY, font=FONT(11 * S), anchor="rm"); d.text((cx + m["l"] - 6 * S, Y(-mx)), f"-{mx:.2f}", fill=GRAY, font=FONT(11 * S), anchor="rm")
    img = img.resize((img.width // S, img.height // S), Image.LANCZOS); img.save(path); return path

def picks_stats(rows, n=5, h="r1"):
    out = []; ct = cb = cs = 1.0
    for r in rows:
        tv = [x[h] for x in r["top"][:n] if x.get(h) is not None]; bv = [x[h] for x in r["bot"][:n] if x.get(h) is not None]
        if not tv or not bv: continue
        mt, mb = sum(tv) / len(tv), sum(bv) / len(bv); ct *= 1 + mt; cb *= 1 + mb; cs *= 1 + (mt - mb)
        out.append({"w": r["w"], "mt": mt, "mb": mb, "sp": mt - mb, "ct": ct - 1, "cb": cb - 1, "cs": cs - 1})
    return out

def chart_picks(path, keys=("eps_rev", "si", "tone", "att", "BULL")):
    """Cumulative long-top-5 / short-bottom-5 spread, weekly rebalance, universe-relative, per signal."""
    lab = PK["long"]["labels"]; series = {k: picks_stats(PK["long"]["picks"][k]) for k in keys if k in PK["long"]["picks"]}
    W, H = 1100 * S, 420 * S; m = {"l": 80 * S, "r": 300 * S, "t": 50 * S, "b": 50 * S}
    img = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(img)
    d.text((m["l"], 10 * S), "Weekly picks — cumulative spread of long top-5 / short bottom-5, equal weight, rebalanced every Friday, universe-relative, no costs", fill=INK, font=FONT(15 * S, True))
    allv = [x["cs"] for s in series.values() for x in s]; lo, hi = min(0, min(allv)), max(0.05, max(allv)); n = max(len(s) for s in series.values())
    X = lambda i: m["l"] + i / (n - 1) * (W - m["l"] - m["r"]); Y = lambda v: m["t"] + (hi - v) / (hi - lo) * (H - m["t"] - m["b"])
    for g in (lo, 0, hi / 2, hi):
        d.line([m["l"], Y(g), W - m["r"], Y(g)], fill=GRAY if g == 0 else GRID, width=S); d.text((m["l"] - 8 * S, Y(g)), pct(g, 0), fill=GRAY, font=FONT(12 * S), anchor="rm")
    palette = {"eps_rev": BLUE, "si": (27, 175, 122), "tone": ORANGE, "att": (74, 58, 167), "BULL": (232, 123, 164)}
    ws = max(series.values(), key=len)
    for i in range(0, n, max(1, n // 10)): d.text((X(i), H - m["b"] + 8 * S), ws[i]["w"][:7], fill=GRAY, font=FONT(12 * S), anchor="ma")
    for j, (k, s) in enumerate(series.items()):
        pts = [(X(i), Y(x["cs"])) for i, x in enumerate(s)]; d.line(pts, fill=palette.get(k, GRAY), width=2 * S, joint="curve")
        d.rectangle([W - m["r"] + 14 * S, m["t"] + j * 26 * S, W - m["r"] + 28 * S, m["t"] + j * 26 * S + 14 * S], fill=palette.get(k, GRAY))
        d.text((W - m["r"] + 36 * S, m["t"] + j * 26 * S + 7 * S), f"{ {'eps_rev':'EPS revision (4w)','si':'Short interest (level)','tone':'Twitter sentiment (BBG)','att':'Twitter attention (BBG)','BULL':'BULL composite'}.get(k, lab.get(k, k))[:30]}  {pct(s[-1]['cs'], 0)}", fill=INK2, font=FONT(13 * S), anchor="lm")
    img = img.resize((img.width // S, img.height // S), Image.LANCZOS); img.save(path); return path

def chart_events(path):
    """Pre-print signals → 15-day relative return (IC), horizontal bars with n."""
    st = EV["stats"]; SIG = EV["signals"]
    rows = [(k, st[f"{k}|tot15"]) for k in SIG if f"{k}|tot15" in st and not k.startswith(("post_", "own_post", "ss_post"))]
    rows.sort(key=lambda x: -x[1]["ic"])
    W, H = 1150 * S, (60 + 30 * len(rows)) * S; m = {"l": 430 * S, "r": 330 * S, "t": 44 * S, "b": 16 * S}
    img = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(img)
    d.text((20 * S, 10 * S), f"Earnings windows — pre-print signal (10 days before) → 15-day relative return after the print, rank IC across {EV['n_events']:,} prints (Jan-2024 → Aug-2026)", fill=INK, font=FONT(15 * S, True))
    mx = max(0.1, *[abs(r[1]["ic"]) for r in rows]); X = lambda v: m["l"] + (v + mx) / (2 * mx) * (W - m["l"] - m["r"]); bh = 30 * S
    d.line([X(0), m["t"], X(0), H - m["b"]], fill=GRAY, width=S)
    for i, (k, s) in enumerate(rows):
        y = m["t"] + i * bh; v = s["ic"]; col = BLUE if v >= 0 else RED
        d.text((m["l"] - 10 * S, y + bh / 2), SIG[k][:58], fill=INK2, font=FONT(13 * S), anchor="rm")
        d.rectangle([min(X(0), X(v)), y + 6 * S, max(X(0), X(v)), y + bh - 6 * S], fill=col)
        d.text((max(X(0), X(v)) + 6 * S, y + bh / 2), f"{v:+.3f}  (n={s['n']}, Q5−Q1 {pct(s['q5_q1'])})", fill=INK2, font=FONT(12 * S), anchor="lm")
    img = img.resize((img.width // S, img.height // S), Image.LANCZOS); img.save(path); return path

# ----------------------------------------------------------------------------- docx helpers
def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_); tcPr.append(shd)

def table(doc, header, rows, widths=None, color_cols=(), size=9):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; p = c.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.size = Pt(size); shade(c, "EEECE7")
        if i > 0: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = cells[i].paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(size)
            if i > 0: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i in color_cols and isinstance(v, str) and v[:1] in "+-" and v not in ("–",):
                r.font.color.rgb = RGBColor(0x2A, 0x78, 0xD6) if v.startswith("+") else RGBColor(0xC0, 0x39, 0x38)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows: row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t

def H(doc, text, lvl=1): doc.add_heading(text, level=lvl)
def P(doc, text, bold_lead=None, size=10.5, italic=False):
    p = doc.add_paragraph()
    if bold_lead: r = p.add_run(bold_lead + " "); r.bold = True; r.font.size = Pt(size)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic; return p
def B(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple): r = p.add_run(it[0] + " "); r.bold = True; r.font.size = Pt(size); it = it[1]
        r = p.add_run(it); r.font.size = Pt(size)

def bt(src, k, h): return src["backtest"].get(f"{k}|rel_ret_{h}", {})

# ----------------------------------------------------------------------------- build
def build(out):
    tmp = Path(__file__).parent / "_rpt_tmp"; tmp.mkdir(exist_ok=True)
    c1 = chart_ccf(tmp / "ccf.png"); c2 = chart_picks(tmp / "picks.png"); c3 = chart_events(tmp / "events.png")
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in doc.sections: s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)

    lw = PW["weeks"]; last = lw[-1]; L = LL["long"]; lab = LL["labels"]
    g = lambda k, l: (L["ccf"][k].get(str(l)) or {}).get("ic")
    par = lambda k: (L["partial"].get(k) or {}).get("1") or {}
    ev = EV["stats"]; e15 = lambda k: ev.get(f"{k}|tot15", {}); er = lambda k: ev.get(f"{k}|react", {})
    pk5 = {k: picks_stats(PK["long"]["picks"][k]) for k in PK["long"]["picks"]}
    pkm = lambda k: (sum(x["mt"] for x in pk5[k]) / len(pk5[k]), sum(x["mb"] for x in pk5[k]) / len(pk5[k]), sum(x["sp"] for x in pk5[k]) / len(pk5[k]), sum(1 for x in pk5[k] if x["sp"] > 0) / len(pk5[k]), pk5[k][-1]["cs"], len(pk5[k]))

    # ---- title
    t = doc.add_heading("Sentiment backtest — does hot/cold × bull/bear lead the tape?", 0)
    P(doc, f"Capstone US Equities · internal · {TODAY.strftime('%d %B %Y')} · universe of {len(PW['tickers'])} wiki names · prepared from the sentiment dashboard build (build_sentiment.py). Every figure below is reproducible from _wiki/_data/sentiment/.", size=9.5, italic=True)

    # ---- verdict
    H(doc, "1. Bottom line")
    P(doc, "Sentiment is coincident to lagging, not leading. Twitter and news tone move with the same week's return and after it; they carry almost no information about the following week. "
           "The only signals with forward information in this universe are analyst EPS revisions (a small, well-known residual that survives a momentum control) and short interest as a slow positioning level. "
           "Sentiment does earn a place in one setting: positioning into a known catalyst. Pre-print tone and attention predict the earnings reaction and the 15 days after it, and being 'hot' into a print raises the bar for the result.")
    B(doc, [
        ("Tone is coincident.", f"Bloomberg Twitter sentiment: rank IC {sgn(g('tone', 0))} with the same week's relative return, {sgn(g('tone', -1))} with the prior week, {sgn(g('tone', 1))} with the next week. News sentiment: {sgn(g('news_tone', 0))} same week. The reverse test (this week's return → next week's tone) is {sgn(L['reverse']['tone']['ic'])}: price moves first, the crowd talks afterwards. Identical shape in 2024, 2025 and 2026."),
        ("Attention is a fade, not a lead.", f"All-of-X tweet volume vs its own trailing month: IC {sgn(g('att', 0))} same week, {sgn(g('att', 1))} next week. As a weekly top-5 basket the hottest names returned {pct(pkm('att')[0], 2)} the following week and the coldest {pct(pkm('att')[1], 2)}."),
        ("EPS revisions are the one residual.", f"Mostly lagging (IC ≈ +0.10 on the four prior weeks: revisions chase price) but with a genuine forward tail: {sgn(g('eps_rev', 1))} next week, {sgn(par('eps_rev').get('ic_partial'))} after removing the past-4-week return (t {fmt(par('eps_rev').get('t_partial'))}), persisting to four weeks out. Top-5 basket {pct(pkm('eps_rev')[0], 2)}/week vs bottom-5 {pct(pkm('eps_rev')[1], 2)}, spread positive in {pkm('eps_rev')[3]:.0%} of weeks, {pct(pkm('eps_rev')[4], 0)} cumulative over {pkm('eps_rev')[5]} weeks."),
        ("Short interest is a level signal.", f"Flat week-to-week correlations, but inside every momentum tercile the high-SI third beat the low-SI third by +1.3% to +1.8% over 15 days (2025–26 squeeze regime). Top-5 basket {pct(pkm('si')[0], 2)}/week, spread {pct(pkm('si')[2], 2)}, {pct(pkm('si')[4], 0)} cumulative."),
        ("Where sentiment does work: into earnings.", f"Across {EV['n_events']:,} prints, pre-print Twitter sentiment has IC {sgn(e15('bbg_tone').get('ic'))} on the 15-day return (top-minus-bottom quintile {pct(e15('bbg_tone').get('q5_q1'))}); news sentiment {sgn(er('news_tone').get('ic'))} on the reaction. Hot-into-print names that beat earned {pct(EV['conditional'].get('hot|beat', {}).get('tot15'))} over 15 days vs {pct(EV['conditional'].get('cold|beat', {}).get('tot15'))} for cold names that beat; hot names that printed in line lost {pct(EV['conditional'].get('hot|inline', {}).get('tot15'))}. Post-print sentiment shifts confirm the move (IC {sgn(ev.get('post_tone_shift|react', {}).get('ic'))} on the reaction) and do not predict the drift ({sgn(ev.get('post_tone_shift|drift', {}).get('ic'))})."),
        ("Own corpus and sell-side: too short to conclude.", f"The curated-handle corpus (May-2026→, {len(lw)} weeks) and the Outlook sell-side leg (Feb-2026→) show nothing with |t| ≥ 2 on the weekly test. An earlier read that curated attention leads (15-day IC +0.04) does not survive the non-overlapping weekly test and should be treated as noise. Sell-side note flow is contrarian; pre-print broker upgrades are the most contrarian pre-print signal in the study (IC {sgn(er('ss_net').get('ic'))} on the reaction, n={er('ss_net').get('n')})."),
    ])
    P(doc, "Practical implication: keep the hot/cold ranking as a read on where the crowd is, not as a buy list. Put forecasting weight on revisions, short interest and the pre-print setup; show tone as context.", bold_lead="")

    # ---- data
    H(doc, "2. Data and construction")
    table(doc, ["Source", "Coverage", "What it contributes"], [
        ["Own tweet corpus (twitter-briefing routine, twitterapi.io)", f"{TW['tweets_scanned']:,} tweets scanned since {TW['start']}, {TW['tweets_matched']:,} mapped to a ticker; ~440 curated handles", "Curated attention (mentions vs own trailing 4 weeks, wires separate) and lexicon tone"],
        ["Bloomberg social/news fields (bdh)", f"Daily Twitter sentiment, tweet count, pos/neg counts, news sentiment and count, all 99 names, Dec-2023 → today", "All-of-X attention and tone — the long backward track (139 weeks)"],
        ["Bloomberg estimates/positioning", "BEST_EPS (1BF) daily, BEST_ANALYST_RATING, SI_PERCENT_EQUITY_FLOAT, PX_LAST", "EPS revision (4-week %), rating drift, short interest level, returns"],
        ["Outlook sell-side e-mails (COM, DASL sweep)", f"{SS.get('emails', 0):,} broker e-mails since {SS.get('since')}, {SS.get('matched', 0):,} mapped to 97 tickers (subject + body)", "Note flow, note tone, rating changes, price-target changes, estimate direction"],
        ["Bloomberg earnings history (bds)", f"{EV['n_events']:,} prints with a full 15-day window, Jan-2024 → Aug-2026, time-of-day aware", "Event windows: reaction, 15-day total, drift; EPS surprise"],
    ], widths=[2.1, 2.4, 2.4])
    B(doc, [
        ("Construction.", "Signals are stamped every Friday and z-scored across names within the week (clipped ±3). Attention = log((count this week + 0.5) / (trailing 4-week mean + 0.5)). Tone = count- or engagement-weighted weekly mean, shrunk toward zero when few observations. HOT = mean z of the attention legs; BULL = mean z of tone, tone change, EPS revision, rating drift (plus sell-side tone and PT change on the own-corpus track)."),
        ("Returns.", "Friday-to-Friday closes in local currency; every test uses the return relative to the equal-weight universe, so market beta is removed. Forward horizons 1 week, 15 calendar days, 4 weeks."),
        ("Tests.", "(a) Weekly rank IC (Spearman across names) of each signal vs forward relative return, with quintile spreads and by-year splits. (b) Cross-correlation function: IC of the signal at week t with the return of week t+k, k = −4…+4, on non-overlapping weekly returns; partial IC controlling for the past-4-week return; tercile double sorts. (c) Weekly top/bottom-N baskets and their next-week return. (d) Earnings event windows: 10-day pre-print signals vs reaction, 15-day total and drift."),
        ("Caveats.", "Forward windows longer than a week overlap, so their t-stats overstate independence; the lead/lag section uses non-overlapping weekly returns for that reason. Lexicon tone is not LLM-scored. The own corpus and sell-side legs cover only 2026. Non-USD names carry FX noise in raw returns."),
    ])

    # ---- weekly backtest
    H(doc, "3. Weekly panel — forward IC by horizon")
    P(doc, f"Bloomberg track, {PL['weeks'][0]} → {PL['weeks'][-1]} ({len(PL['weeks'])} weekly cross-sections). Positive IC = high score preceded outperformance. Q5−Q1 = top minus bottom quintile mean forward relative return.", size=9.5, italic=True)
    rows = []
    for k in list(PL["components"]) + ["HOT", "BULL"]:
        a, b, c = bt(PL, k, "1w"), bt(PL, k, "15d"), bt(PL, k, "4w"); by = c.get("by_year", {})
        rows.append([PL["components"].get(k, [k])[0] if k in PL["components"] else k + " composite", sgn(a.get("mean_ic")), sgn(b.get("mean_ic")), sgn(c.get("mean_ic")), fmt(c.get("t")), f"{c.get('hit', 0):.0%}" if c else "–", pct(c.get("q5_q1")),
                     " / ".join(f"{y[2:]}:{v['mean_ic']:+.3f}" for y, v in by.items())])
    table(doc, ["Signal", "IC 1w", "IC 15d", "IC 4w", "t (4w)", "hit", "Q5−Q1 4w", "IC 4w by year"], rows, widths=[2.2, 0.6, 0.6, 0.6, 0.5, 0.5, 0.75, 1.35], color_cols=(1, 2, 3, 6), size=8.5)
    P(doc, f"Own corpus + sell-side, {lw[0]} → {last} ({len(lw)} weeks — indicative only).", size=9.5, italic=True)
    rows = []
    for k in list(PW["components"]) + ["HOT", "BULL"]:
        a, b, c = bt(PW, k, "1w"), bt(PW, k, "15d"), bt(PW, k, "4w")
        if not b: continue
        rows.append([PW["components"].get(k, [k])[0] if k in PW["components"] else k + " composite", sgn(a.get("mean_ic")), sgn(b.get("mean_ic")), fmt(b.get("t")), f"{b.get('hit', 0):.0%}", pct(b.get("q5_q1")), sgn(c.get("mean_ic")), str(b.get("weeks"))])
    table(doc, ["Signal", "IC 1w", "IC 15d", "t (15d)", "hit", "Q5−Q1 15d", "IC 4w", "weeks"], rows, widths=[2.6, 0.6, 0.6, 0.55, 0.5, 0.8, 0.6, 0.5], color_cols=(1, 2, 5, 6), size=8.5)

    # ---- lead/lag
    H(doc, "4. Lead or lag? Cross-correlation with weekly returns")
    P(doc, "The forward-IC tables cannot separate a leading signal from sentiment that follows price. The cross-correlation function does: for a signal measured at the end of week t, the bars show its rank correlation with the relative return of week t+k. Bars left of the shaded column are returns that happened before the signal (sentiment reacting to price); bars to the right are returns that came after (the signal leading). Weekly returns do not overlap, so each bar's t-stat is honest.")
    doc.add_picture(str(c1), width=Inches(6.7))
    rows = []
    for k in L["ccf"]:
        p_ = par(k)
        rows.append([lab.get(k, k), sgn(g(k, -2)), sgn(g(k, -1)), sgn(g(k, 0)), sgn(g(k, 1)), sgn(g(k, 2)), sgn(p_.get("ic_partial")), fmt(p_.get("t_partial")), L["verdict"].get(k, "")])
    table(doc, ["Signal", "k=−2", "k=−1", "k=0", "k=+1", "k=+2", "partial +1 (ex-momentum)", "t", "Verdict"], rows, widths=[1.9, 0.5, 0.5, 0.5, 0.5, 0.5, 0.8, 0.4, 1.4], color_cols=(1, 2, 3, 4, 5, 6), size=8)
    P(doc, "Reverse direction — this week's return → next week's signal. A large positive number is the signature of a lagging signal.", size=9.5, italic=True)
    table(doc, ["Signal", "IC", "t", "weeks"], [[lab.get(k, k), sgn(v["ic"]), fmt(v["t"]), str(v["n"])] for k, v in L["reverse"].items()], widths=[3.2, 0.8, 0.6, 0.6], color_cols=(1,), size=9)
    P(doc, "Double sort — tercile of past-4-week return (rows) × tercile of the signal (columns) → mean forward 15-day relative return. If a signal only proxies momentum the columns look alike within each row; if it adds information the right column beats the left inside every row.", size=9.5, italic=True)
    for k in ("eps_rev", "si", "tone", "att"):
        c = L["double_sort"].get(k, {})
        if not c: continue
        rows = []
        for r_ in ("mom_lo", "mom_mid", "mom_hi"):
            v = [(c.get(f"{r_}|{s_}") or {}).get("mean") for s_ in ("sig_lo", "sig_mid", "sig_hi")]
            rows.append([r_.replace("mom_", "past 4w "), pct(v[0]), pct(v[1]), pct(v[2]), pct((v[2] - v[0]) if None not in (v[0], v[2]) else None)])
        P(doc, lab.get(k, k), bold_lead="", size=9.5)
        table(doc, ["", "signal low", "signal mid", "signal high", "high − low"], rows, widths=[1.4, 1.0, 1.0, 1.0, 1.0], color_cols=(1, 2, 3, 4), size=8.5)

    # ---- picks
    H(doc, "5. What we would have seen — weekly top/bottom-5 baskets")
    P(doc, "Every Friday, rank the universe on one signal, take the five highest and five lowest names, and record their relative return over the following week. The chart compounds the long-top / short-bottom spread; the table shows the average pick outcome per signal.")
    doc.add_picture(str(c2), width=Inches(6.7))
    rows = []
    for k in sorted(pk5, key=lambda k: -pkm(k)[2]):
        a, b, s_, hit, cum, n = pkm(k)
        rows.append([PK["long"]["labels"].get(k, k), pct(a, 2), pct(b, 2), pct(s_, 2), f"{hit:.0%}", pct(cum, 0), str(n)])
    table(doc, ["Signal (Bloomberg track)", "top-5 next wk", "bottom-5 next wk", "spread", "spread > 0", "cumulative", "weeks"], rows, widths=[2.4, 0.85, 0.95, 0.7, 0.7, 0.8, 0.5], color_cols=(1, 2, 3, 5), size=8.5)
    pko = {k: picks_stats(PK["own"]["picks"][k]) for k in PK["own"]["picks"]}
    rows = []
    for k in sorted(pko, key=lambda k: -(sum(x["sp"] for x in pko[k]) / len(pko[k]))):
        s = pko[k]; a = sum(x["mt"] for x in s) / len(s); b = sum(x["mb"] for x in s) / len(s)
        rows.append([PK["own"]["labels"].get(k, k), pct(a, 2), pct(b, 2), pct(a - b, 2), f"{sum(1 for x in s if x['sp'] > 0) / len(s):.0%}", pct(s[-1]["cs"], 0), str(len(s))])
    P(doc, f"Own corpus + sell-side track, {PK['own']['span'][0]} → {PK['own']['span'][1]} — indicative only.", size=9.5, italic=True)
    table(doc, ["Signal (own track)", "top-5 next wk", "bottom-5 next wk", "spread", "spread > 0", "cumulative", "weeks"], rows, widths=[2.4, 0.85, 0.95, 0.7, 0.7, 0.8, 0.5], color_cols=(1, 2, 3, 5), size=8.5)
    P(doc, "Reading the tone row: the spread is positive because the most-criticised names kept falling, not because the most-praised names rose. Tone is a warning on the losers, not a buy list. The attention row is the mirror image of a lead: the hottest names underperformed the following week.", size=10)

    # ---- events
    H(doc, "6. Earnings event windows")
    P(doc, f"For every print of the 99 names ({EV['n_events']:,} events with a full 15-day window): signals measured over the 10 calendar days before the print (baseline = the 30 days before that); outcomes from the last close before the print — reaction (to the next close), 15 days total, drift (post-print close to day +15). All relative to the equal-weight universe. Own-corpus signals exist from {EV.get('own_from')}, sell-side from {EV.get('ss_from')}.")
    doc.add_picture(str(c3), width=Inches(6.7))
    rows = []
    for k, name in EV["signals"].items():
        r_, t_, d_ = ev.get(f"{k}|react", {}), ev.get(f"{k}|tot15", {}), ev.get(f"{k}|drift", {})
        if not (r_ or t_): continue
        rows.append([name, sgn(r_.get("ic")), pct(r_.get("q5_q1")), sgn(t_.get("ic")), pct(t_.get("q5_q1")), sgn(d_.get("ic")), str(t_.get("n") or r_.get("n"))])
    table(doc, ["Signal", "IC reaction", "Q5−Q1", "IC 15d", "Q5−Q1", "IC drift", "n"], rows, widths=[2.7, 0.7, 0.65, 0.65, 0.65, 0.65, 0.45], color_cols=(1, 2, 3, 4, 5), size=8)
    P(doc, "Attention into the print × result (Bloomberg Twitter attention terciles; beat/miss = EPS surprise beyond ±2%). Read across a row to see whether being hot into the print changes the payoff for the same result.", size=9.5, italic=True)
    C = EV["conditional"]; rows = []
    for a in ("hot", "mid", "cold"):
        r_ = [a]
        for b in ("beat", "inline", "miss"):
            v = C.get(f"{a}|{b}", {}); r_ += [pct(v.get("react")), pct(v.get("tot15")), str(v.get("n", 0))]
        rows.append(r_)
    table(doc, ["attention", "beat: react", "15d", "n", "inline: react", "15d", "n", "miss: react", "15d", "n"], rows, widths=[0.8, 0.7, 0.6, 0.4, 0.7, 0.6, 0.4, 0.7, 0.6, 0.4], color_cols=(1, 2, 4, 5, 7, 8), size=8.5)

    # ---- implications
    H(doc, "7. Implications and next steps")
    B(doc, [
        ("Reframe the indicator.", "Present hot/cold as a crowding and positioning read (who is being talked about, who is being upgraded into a print), and bull/bear as context. Neither is a forecast of next week's tape."),
        ("Put the forecasting weight where the evidence is.", "EPS revision momentum (small, persistent, fading in 2026), short interest as a level (squeeze regime), and the pre-print setup: sentiment and curated attention in the 10 days before a print, crossed with the result."),
        ("Watch the hot-into-print penalty.", "Hot names that beat earned less than a quarter of what cold names that beat earned over 15 days; hot names that printed in line lost 5%. That is an actionable positioning rule ahead of the season."),
        ("Let the own-corpus and sell-side tracks accumulate.", "Eighteen weeks is not a sample. The nightly sweep now runs (17:40 weekdays) and both tracks rebuild in the dashboard; revisit after a full earnings season."),
        ("Upgrades worth testing.", "LLM-scored tone on the curated tweets and the sell-side bodies (the lexicon is the weakest link); an event-window version of the sell-side leg with bodies; options skew as a second crowding leg."),
    ])
    P(doc, "Files: _wiki/_dashboards/sentiment.html (interactive, rebuilt nightly) · _wiki/_tools/build_sentiment.py (steps tweets, bbg, bbglong, mail_merge, sellside, panel, long, leadlag, picks, events_pull, events, dash) · this report: _wiki/_tools/build_sentiment_report.py.", size=9, italic=True)
    doc.save(out); print("wrote", out)
    return out

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default=str(DATA / f"Sentiment_Backtest_{TODAY.isoformat()}.docx"))
    a = ap.parse_args(); build(Path(a.out))
