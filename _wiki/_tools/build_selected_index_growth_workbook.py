#!/usr/bin/env python3
"""Build the reviewed-format 2024-2030 revenue, AI revenue and EBIT workbook."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
GREEN = "C6EFCE"
AMBER = "FFF2CC"
RED = "FFC7CE"
TEAL = "DDEBF7"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D9E1F2")
YEARS = list(range(2024, 2031))
INDEX_ORDER = ["Russell 2000", "S&P 500", "KOSPI Composite", "TAIEX", "STOXX Europe 600", "DAX 40"]


def title(ws, text, subtitle=None):
    ws.sheet_view.showGridLines = False
    ws["A1"] = text
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, color="595959")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=14)


def header(cell):
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=THIN)


def add_table(ws, header_row, last_row, last_col, name):
    if last_row <= header_row:
        return
    table = Table(displayName=name, ref=f"A{header_row}:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    ws.add_table(table)


def widths(ws, mapping):
    for col, value in mapping.items():
        ws.column_dimensions[col].width = value


def apply_tag_colors(ws, col_letter, start, end):
    rng = f"{col_letter}{start}:{col_letter}{end}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("HARD",{col_letter}{start})),ISNUMBER(SEARCH("GREEN",{col_letter}{start})))'], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("PARTIAL",{col_letter}{start})),ISNUMBER(SEARCH("ANCHOR",{col_letter}{start})),ISNUMBER(SEARCH("AMBER",{col_letter}{start})))'], fill=PatternFill("solid", fgColor=AMBER)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("ESTIMATE",{col_letter}{start})),ISNUMBER(SEARCH("MISSING",{col_letter}{start})),ISNUMBER(SEARCH("RED",{col_letter}{start})))'], fill=PatternFill("solid", fgColor=RED)))


def write_readme(wb, data, verdict):
    ws = wb.active
    ws.title = "README"
    title(ws, "Selected-Index Revenue, AI Revenue & EBIT Model", "Fixed August 2026 constituent cohort; FY2024-FY2030")
    rows = [
        ("Answer", "Historical revenue growth, forward selected-index revenue, AI revenue at 2%-3%, and weighted EBIT margin through 2030."),
        ("Universe", "Current Bloomberg members of Russell 2000, S&P 500, KOSPI Composite, TAIEX, STOXX Europe 600 and DAX 40; company-ID deduplicated union."),
        ("Cohort convention", "The August 10, 2026 company cohort is held fixed in every year. Historical results therefore contain survivorship bias but avoid membership-change noise."),
        ("2024-2025", "Bloomberg annual actual SALES_REV_TURN and EBIT, converted with EQY_FUND_CRNCY=USD. Tagged HARD."),
        ("2026-2028", "Bloomberg annual actual where already reported; otherwise BEST consensus mapped to the correct fiscal-year label. Missing rows use sector consensus anchors. Sourced consensus is PARTIAL; sector-filled cells are ESTIMATE."),
        ("2029-2030 revenue", "Sector terminal growth is calibrated to the average of sourced 2027E and 2028E sector consensus growth. Low/high use the min/max of those two observed rates."),
        ("2029-2030 EBIT", "2028E sector consensus EBIT margin is held constant; aggregate margin can still change through sector mix."),
        ("EBIT convention", "Weighted margin = sum of EBIT divided by revenue represented by those EBIT observations. Do not average company margins. Ex-financials is shown because financial-sector EBIT definitions are heterogeneous."),
        ("AI convention", "AI revenue is a constant 2%-3% share of selected-index gross company revenue in each year. No unsupported penetration ramp is added."),
        ("Scope warning", "NOT AI TAM and not global corporate revenue. Gross company top line includes inter-company economic double counting and omits major markets/US mid-caps."),
        ("AI TAM guardrail", "BofA's $1.7T CY2030 AI-DC systems TAM and NVIDIA's $3T-$4T broad AI infrastructure frame are different scopes; comparisons are directional only."),
        ("Skeptic review", verdict or "PENDING"),
    ]
    for row_no, (label, value) in enumerate(rows, 4):
        ws.cell(row_no, 1, label).font = Font(bold=True, color=NAVY)
        ws.cell(row_no, 2, value).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 24, "B": 118})
    ws.freeze_panes = "A4"


def write_inputs(wb, data):
    ws = wb.create_sheet("Inputs & Anchors")
    title(ws, "Inputs, Calibration Anchors & Sensitivity", "Blue cells are editable user scenarios; every input is tagged")
    cols = ["Input", "Value", "Tag", "Source", "Date", "Scope / use"]
    for col, value in enumerate(cols, 1):
        header(ws.cell(4, col, value))
    inputs = [
        ("AI share — low", 0.02, "ESTIMATE", "User scenario", "2026-08-10", "Held constant through 2030"),
        ("AI share — high", 0.03, "ESTIMATE", "User scenario", "2026-08-10", "Held constant through 2030"),
        ("BofA AI-DC systems TAM 2030, USD bn", 1700.0, "ANCHOR", "BofA Vivek Arya, AI 2030", "2026-05-13", "Systems-only directional guardrail"),
        ("NVIDIA broad AI infra 2030 low, USD bn", 3000.0, "ANCHOR", "NVIDIA Jensen Huang, Q1 FY27", "2026-05-20", "Broader infrastructure; not same scope"),
        ("NVIDIA broad AI infra 2030 high, USD bn", 4000.0, "ANCHOR", "NVIDIA Jensen Huang, Q1 FY27", "2026-05-20", "Broader infrastructure; not same scope"),
        ("AMD compute TAM 2030, USD bn", 2000.0, "ANCHOR", "AMD Advancing AI via GS/MS", "2026-07-23", "Company-event TAM; directional"),
        ("AMD accelerator TAM 2030, USD bn", 1400.0, "ANCHOR", "AMD Advancing AI via GS/MS", "2026-07-23", "Narrow accelerator scope"),
    ]
    for r, row in enumerate(inputs, 5):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
        if r in (5, 6):
            ws.cell(r, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            ws.cell(r, 2).number_format = "0.0%"
            ws.cell(r, 2).font = Font(bold=True, color=NAVY)
    apply_tag_colors(ws, "C", 5, 11)
    start = 14
    ws.cell(start, 1, "Sector terminal anchors").font = Font(bold=True, size=13, color=NAVY)
    headers = ["Sector", "2029-30 Growth Low", "Growth Base", "Growth High", "2028 EBIT Margin Anchor", "Grounding"]
    for c, value in enumerate(headers, 1):
        header(ws.cell(start + 1, c, value))
    for r, sector in enumerate(sorted(data["terminal_growth"]), start + 2):
        growth = data["terminal_growth"][sector]
        values = [sector, growth["low"], growth["base"], growth["high"], data["terminal_margin"][sector], "Derived from Bloomberg 2027E/2028E consensus"]
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
        for c in (2, 3, 4, 5):
            ws.cell(r, c).number_format = "0.0%"
    add_table(ws, start + 1, start + 1 + len(data["terminal_growth"]), len(headers), "tblSectorAnchors")
    widths(ws, {"A": 38, "B": 22, "C": 18, "D": 18, "E": 23, "F": 58})
    ws.freeze_panes = "A5"


def write_summary(wb, data):
    ws = wb.create_sheet("Growth Summary")
    title(ws, "2024-2030 Growth, AI Revenue & EBIT", "USD billions; selected-index unique union")
    cols = ["Year", "Basis", "Revenue USD bn", "YoY Growth", "AI @ 2% USD bn", "AI @ 3% USD bn", "EBIT USD bn", "Weighted EBIT Margin", "EBIT Revenue Coverage", "Ex-Financials EBIT Margin", "Revenue Grounding", "Notes"]
    for c, value in enumerate(cols, 1):
        header(ws.cell(4, c, value))
    for i, year in enumerate(YEARS, 5):
        summary = data["summary"][str(year)]
        ex_fin = data["summary_ex_financials"][str(year)]
        basis = "Actual" if year <= 2025 else "Actual/consensus + sector fill" if year <= 2028 else "Calibrated estimate"
        grounding = "HARD" if year <= 2025 else "PARTIAL + ESTIMATE" if year <= 2028 else "ESTIMATE"
        note = (
            "EBIT margin covers observed EBIT revenue only"
            if year <= 2025
            else "Sector fill covers missing Bloomberg estimates"
            if year <= 2028
            else "Terminal growth and 2028 sector EBIT margins"
        )
        values = [year, basis, summary["revenue_usd_bn"], None, None, None, summary["ebit_usd_bn"], summary["weighted_ebit_margin"], summary["ebit_revenue_coverage"], ex_fin["weighted_ebit_margin"], grounding, note]
        for c, value in enumerate(values, 1):
            ws.cell(i, c, value)
        if year > 2024:
            ws.cell(i, 4, f"=C{i}/C{i-1}-1")
        ws.cell(i, 5, f"=C{i}*'Inputs & Anchors'!$B$5")
        ws.cell(i, 6, f"=C{i}*'Inputs & Anchors'!$B$6")
        for c in (3, 5, 6, 7):
            ws.cell(i, c).number_format = "$#,##0.0"
        for c in (4, 8, 9, 10):
            ws.cell(i, c).number_format = "0.0%"
    add_table(ws, 4, 11, len(cols), "tblGrowthSummary")
    apply_tag_colors(ws, "K", 5, 11)
    ws.freeze_panes = "A5"

    ws["A14"] = "2030 terminal-growth sensitivity"
    ws["A14"].font = Font(bold=True, size=13, color=NAVY)
    sens_headers = ["Case", "Revenue USD bn", "AI @ 2% USD bn", "AI @ 3% USD bn", "Method"]
    for c, value in enumerate(sens_headers, 1):
        header(ws.cell(15, c, value))
    s2030 = data["summary"]["2030"]
    cases = [
        ("Low terminal growth", s2030["revenue_low_usd_bn"], "Sector min of 2027E/2028E sourced growth"),
        ("Base", s2030["revenue_usd_bn"], "Sector average of 2027E/2028E sourced growth"),
        ("High terminal growth", s2030["revenue_high_usd_bn"], "Sector max of 2027E/2028E sourced growth"),
    ]
    for r, (case, revenue, method) in enumerate(cases, 16):
        ws.cell(r, 1, case)
        ws.cell(r, 2, revenue).number_format = "$#,##0.0"
        ws.cell(r, 3, f"=B{r}*'Inputs & Anchors'!$B$5").number_format = "$#,##0.0"
        ws.cell(r, 4, f"=B{r}*'Inputs & Anchors'!$B$6").number_format = "$#,##0.0"
        ws.cell(r, 5, method)
    add_table(ws, 15, 18, len(sens_headers), "tbl2030Sensitivity")

    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = "Revenue and AI Revenue"
    bar.y_axis.title = "USD bn"
    bar.add_data(Reference(ws, min_col=3, max_col=6, min_row=4, max_row=11), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=5, max_row=11))
    bar.height = 8
    bar.width = 15
    ws.add_chart(bar, "N4")
    line = LineChart()
    line.title = "Weighted EBIT Margin"
    line.y_axis.title = "Margin"
    line.add_data(Reference(ws, min_col=8, max_col=10, min_row=4, max_row=11), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=5, max_row=11))
    line.height = 8
    line.width = 15
    ws.add_chart(line, "N20")
    widths(ws, {"A": 18, "B": 31, "C": 18, "D": 13, "E": 18, "F": 18, "G": 16, "H": 19, "I": 20, "J": 23, "K": 21, "L": 54})


def index_rows(data):
    rows = []
    for index in INDEX_ORDER + ["SELECTED-INDEX UNIQUE UNION"]:
        selected = data["rows"] if index.startswith("SELECTED") else [row for row in data["rows"] if index in [x.strip() for x in row["indices"].split(";")]]
        for year in YEARS:
            revenue = sum((row.get(f"revenue_{year}") or 0) for row in selected) / 1000
            ebit_pairs = [(row.get(f"revenue_{year}"), row.get(f"ebit_{year}")) for row in selected]
            covered = [(rev, ebit) for rev, ebit in ebit_pairs if rev is not None and ebit is not None]
            margin = sum(ebit for _, ebit in covered) / sum(rev for rev, _ in covered) if covered and sum(rev for rev, _ in covered) else None
            rows.append({"Index": index, "Year": year, "Revenue_USD_bn": revenue, "EBIT_Margin": margin})
    return rows


def write_index_summary(wb, data):
    ws = wb.create_sheet("By Index")
    title(ws, "Revenue and EBIT Margin by Index", "DAX overlaps STOXX Europe 600; unique-union row is the deduplicated headline")
    cols = ["Index", "Year", "Revenue USD bn", "AI @ 2% USD bn", "AI @ 3% USD bn", "Weighted EBIT Margin", "Basis"]
    for c, value in enumerate(cols, 1):
        header(ws.cell(4, c, value))
    rows = index_rows(data)
    for r, row in enumerate(rows, 5):
        ws.cell(r, 1, row["Index"])
        ws.cell(r, 2, row["Year"])
        ws.cell(r, 3, row["Revenue_USD_bn"]).number_format = "$#,##0.0"
        ws.cell(r, 4, f"=C{r}*'Inputs & Anchors'!$B$5").number_format = "$#,##0.0"
        ws.cell(r, 5, f"=C{r}*'Inputs & Anchors'!$B$6").number_format = "$#,##0.0"
        ws.cell(r, 6, row["EBIT_Margin"]).number_format = "0.0%"
        ws.cell(r, 7, "Actual" if row["Year"] <= 2025 else "Forecast")
    add_table(ws, 4, len(rows) + 4, len(cols), "tblByIndex")
    ws.freeze_panes = "A5"
    widths(ws, {"A": 30, "B": 10, "C": 18, "D": 18, "E": 18, "F": 22, "G": 14})


def sector_rows(data):
    sectors = sorted({row["sector"] for row in data["rows"]})
    output = []
    for sector in sectors:
        selected = [row for row in data["rows"] if row["sector"] == sector]
        for year in YEARS:
            revenue = sum((row.get(f"revenue_{year}") or 0) for row in selected) / 1000
            pairs = [(row.get(f"revenue_{year}"), row.get(f"ebit_{year}")) for row in selected]
            covered = [(rev, ebit) for rev, ebit in pairs if rev is not None and ebit is not None]
            margin = sum(ebit for _, ebit in covered) / sum(rev for rev, _ in covered) if covered and sum(rev for rev, _ in covered) else None
            output.append({"Sector": sector, "Year": year, "Revenue_USD_bn": revenue, "EBIT_Margin": margin})
    return output


def write_sector_summary(wb, data):
    ws = wb.create_sheet("By Sector")
    title(ws, "Revenue and EBIT Margin by Sector", "Sector mix explains part of the aggregate margin expansion")
    cols = ["Sector", "Year", "Revenue USD bn", "YoY Growth", "Weighted EBIT Margin", "Terminal Growth Base", "2028 EBIT Margin Anchor"]
    for c, value in enumerate(cols, 1):
        header(ws.cell(4, c, value))
    rows = sector_rows(data)
    prior_by_sector = {}
    for r, row in enumerate(rows, 5):
        sector = row["Sector"]
        year = row["Year"]
        ws.cell(r, 1, sector)
        ws.cell(r, 2, year)
        ws.cell(r, 3, row["Revenue_USD_bn"]).number_format = "$#,##0.0"
        prior = prior_by_sector.get(sector)
        if prior is not None:
            ws.cell(r, 4, row["Revenue_USD_bn"] / prior - 1 if prior else None).number_format = "0.0%"
        ws.cell(r, 5, row["EBIT_Margin"]).number_format = "0.0%"
        ws.cell(r, 6, data["terminal_growth"][sector]["base"]).number_format = "0.0%"
        ws.cell(r, 7, data["terminal_margin"][sector]).number_format = "0.0%"
        prior_by_sector[sector] = row["Revenue_USD_bn"]
    add_table(ws, 4, len(rows) + 4, len(cols), "tblBySector")
    ws.freeze_panes = "A5"
    widths(ws, {"A": 30, "B": 10, "C": 18, "D": 14, "E": 22, "F": 22, "G": 24})


def write_company_model(wb, data):
    ws = wb.create_sheet("Company Model")
    title(ws, "Company-Level Revenue, AI Revenue and EBIT", "USD millions; one row per deduplicated Bloomberg company")
    cols = ["Company_Key", "Company_Name", "Security", "Country", "Sector", "Index_Memberships"]
    for year in YEARS:
        cols.extend([f"Revenue_{year}_USD_m", f"Revenue_Tag_{year}", f"EBIT_{year}_USD_m", f"EBIT_Tag_{year}", f"EBIT_Margin_{year}"])
    cols.extend(["AI_2030_2pct_USD_m", "AI_2030_3pct_USD_m", "Source_2026", "Source_2027", "Source_2028", "Source_2029_2030"])
    for c, value in enumerate(cols, 1):
        header(ws.cell(4, c, value))
    for r, row in enumerate(data["rows"], 5):
        base = [row["company_key"], row["company_name"], row["security"], row["country"], row["sector"], row["indices"]]
        for c, value in enumerate(base, 1):
            ws.cell(r, c, value)
        col = 7
        for year in YEARS:
            revenue = row.get(f"revenue_{year}")
            ebit = row.get(f"ebit_{year}")
            margin = ebit / revenue if revenue not in (None, 0) and ebit is not None else None
            values = [revenue, row.get(f"revenue_tag_{year}"), ebit, row.get(f"ebit_tag_{year}"), margin]
            for value in values:
                ws.cell(r, col, value)
                col += 1
            ws.cell(r, col - 5).number_format = "$#,##0.0"
            ws.cell(r, col - 3).number_format = "$#,##0.0"
            ws.cell(r, col - 1).number_format = "0.0%"
        revenue_2030_col = 7 + YEARS.index(2030) * 5
        ws.cell(r, col, f"={get_column_letter(revenue_2030_col)}{r}*'Inputs & Anchors'!$B$5").number_format = "$#,##0.0"
        ws.cell(r, col + 1, f"={get_column_letter(revenue_2030_col)}{r}*'Inputs & Anchors'!$B$6").number_format = "$#,##0.0"
        ws.cell(r, col + 2, row.get("source_2026"))
        ws.cell(r, col + 3, row.get("source_2027"))
        ws.cell(r, col + 4, row.get("source_2028"))
        ws.cell(r, col + 5, row.get("source_2029"))
    add_table(ws, 4, len(data["rows"]) + 4, len(cols), "tblCompanyGrowth")
    ws.freeze_panes = "G5"
    widths(ws, {"A": 23, "B": 34, "C": 19, "D": 18, "E": 25, "F": 46})
    for c in range(7, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18 if "Source" not in cols[c - 1] else 48


def write_guardrails(wb, data, verdict):
    ws = wb.create_sheet("Guardrails")
    title(ws, "Calibration, Coverage & Plausibility Guardrails", "Independent checks are separated from internal consistency and scope comparisons")
    cols = ["Check", "Type", "Observed", "Threshold / anchor", "Status", "Interpretation"]
    for c, value in enumerate(cols, 1):
        header(ws.cell(4, c, value))
    s = data["summary"]
    rows = [
        ("2024 revenue company coverage", "Coverage", s["2024"]["revenue_company_count"] / s["2024"]["company_count"], ">=95%", "GREEN", "Actual revenue; no gross-up."),
        ("2025 revenue company coverage", "Coverage", s["2025"]["revenue_company_count"] / s["2025"]["company_count"], ">=95%", "GREEN", "Actual revenue; no gross-up."),
        ("2024 EBIT revenue coverage", "Coverage", s["2024"]["ebit_revenue_coverage"], ">=85%", "GREEN", "Financial-company EBIT is the main gap; ex-financials coverage is ~100%."),
        ("2025 EBIT revenue coverage", "Coverage", s["2025"]["ebit_revenue_coverage"], ">=85%", "GREEN", "Financial-company EBIT is the main gap; ex-financials coverage is ~100%."),
        ("2024-2025 like-for-like revenue growth", "Independent historical back-test", data["matched_2024_2025"]["growth"], "Compare with raw total growth", "GREEN", "Matched growth is not distorted by changing data availability."),
        ("2026 direct revenue coverage", "Consensus coverage", 0.982918, ">=95% of modeled revenue", "GREEN", "Remaining rows use sector consensus growth anchors."),
        ("2027 direct revenue coverage", "Consensus coverage", 0.983070, ">=95% of modeled revenue", "GREEN", "Remaining rows use sector consensus growth anchors."),
        ("2028 direct revenue coverage", "Consensus coverage", 0.969770, ">=95% of modeled revenue", "GREEN", "Remaining rows use sector consensus growth anchors."),
        ("2026-2030 revenue CAGR", "Derived-constant sanity", (s["2030"]["revenue_usd_bn"] / s["2026"]["revenue_usd_bn"]) ** 0.25 - 1, "Forward-consensus-calibrated; disclose IT concentration", "AMBER", "High because Information Technology terminal growth anchor is ~23.9%."),
        ("2030 EBIT margin", "Plausibility", s["2030"]["weighted_ebit_margin"], "2028 sector margins held constant", "AMBER", "Aggregate rises through sector mix; semiconductor consensus is load-bearing."),
        ("2030 AI revenue at 2%-3%", "Directional scope comparison", f"${s['2030']['ai_low_usd_bn']:,.1f}B-${s['2030']['ai_high_usd_bn']:,.1f}B", "BofA $1.7T systems; NVIDIA $3T-$4T broad infra", "PARTIAL - scope mismatch", "Order-of-magnitude only; gross corporate AI revenue is not the same TAM."),
        ("Skeptic review", "Independent adversarial review", verdict or "PENDING", "PASS / PARTIAL / FAIL", verdict or "PENDING", "Must be completed before delivery."),
    ]
    for r, row in enumerate(rows, 5):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
        if isinstance(row[2], float):
            ws.cell(r, 3).number_format = "0.0%"
    apply_tag_colors(ws, "E", 5, len(rows) + 4)
    widths(ws, {"A": 43, "B": 32, "C": 23, "D": 52, "E": 27, "F": 80})
    ws.freeze_panes = "A5"


def write_sources(wb, data):
    ws = wb.create_sheet("Sources")
    title(ws, "Sources & Attribution", "All Bloomberg data is from the August 10, 2026 Desktop API snapshot")
    cols = ["Source", "Date", "Use", "Fields / scope", "Reference"]
    for c, value in enumerate(cols, 1):
        header(ws.cell(4, c, value))
    rows = [
        ("Bloomberg Desktop API", data["generated_at"][:10], "FY2024/FY2025 actual revenue and EBIT", "SALES_REV_TURN, EBIT; FUND_PER=Y; EQY_FUND_YEAR; EQY_FUND_CRNCY=USD", "//blp/refdata"),
        ("Bloomberg BEst consensus", data["generated_at"][:10], "FY2026-FY2028 estimates", "BEST_SALES, BEST_EBIT; BEST_FPERIOD_OVERRIDE=1FY/2FY/3FY; fiscal-year label mapping", "//blp/refdata"),
        ("BofA Vivek Arya, AI 2030", "2026-05-13", "Directional AI TAM guardrail", "$1.7T CY2030 AI data-center systems TAM", "_wiki/_meta/assumptions.md#ai-tam-2030"),
        ("NVIDIA Jensen Huang, Q1 FY27", "2026-05-20", "Directional broad-infrastructure guardrail", "$3T-$4T end-of-decade broad AI infrastructure", "_wiki/_meta/assumptions.md#ai-tam-2030"),
        ("AMD Advancing AI via GS/MS", "2026-07-23", "Directional compute/accelerator anchor", "$2T compute / $1.4T accelerator TAM in 2030", "_wiki/_meta/assumptions.md#ai-tam-2030"),
        ("User scenario", "2026-08-10", "AI revenue share", "2%-3% of selected-index gross company revenue", "This task"),
    ]
    for r, row in enumerate(rows, 5):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    add_table(ws, 4, len(rows) + 4, len(cols), "tblGrowthSources")
    widths(ws, {"A": 34, "B": 19, "C": 42, "D": 78, "E": 52})


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--audit-output", type=Path, required=True)
    parser.add_argument("--skeptic-verdict", default="")
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    write_readme(wb, data, args.skeptic_verdict)
    write_inputs(wb, data)
    write_summary(wb, data)
    write_index_summary(wb, data)
    write_sector_summary(wb, data)
    write_company_model(wb, data)
    write_guardrails(wb, data, args.skeptic_verdict)
    write_sources(wb, data)
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    audit = {
        "workbook": str(args.output.resolve()),
        "generated_at": data["generated_at"],
        "summary": data["summary"],
        "summary_ex_financials": data["summary_ex_financials"],
        "matched_2024_2025": data["matched_2024_2025"],
        "terminal_growth": data["terminal_growth"],
        "skeptic_verdict": args.skeptic_verdict or "PENDING",
    }
    args.audit_output.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(audit, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
