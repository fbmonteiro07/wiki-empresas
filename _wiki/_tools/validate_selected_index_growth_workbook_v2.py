#!/usr/bin/env python3
"""Validate the normalized v2 selected-index growth workbook."""

import argparse
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook


EXPECTED = ["README", "Inputs & Anchors", "Growth Summary", "By Index", "By Sector", "Company Model", "Guardrails", "Sources"]


def hmap(ws, row=4):
    return {ws.cell(row, c).value: c for c in range(1, ws.max_column + 1)}


def close(actual, expected, tolerance=1e-8):
    return actual is not None and expected is not None and abs(actual - expected) <= tolerance


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    audit = json.loads(args.audit.read_text(encoding="utf-8"))
    failures = []
    warnings = []
    checks = {}

    with zipfile.ZipFile(args.workbook) as archive:
        bad = archive.testzip()
        checks["xlsx_zip_integrity"] = bad is None
        if bad:
            failures.append(f"Corrupt XLSX member: {bad}")

    wb = load_workbook(args.workbook, data_only=False, read_only=False)
    checks["sheets"] = wb.sheetnames
    if wb.sheetnames != EXPECTED:
        failures.append(f"Unexpected sheet names/order: {wb.sheetnames}")

    summary = wb["Growth Summary"]
    sh = hmap(summary)
    required_summary = ["Year", "Revenue USD bn", "AI @ 2% USD bn", "AI @ 3% USD bn", "Headline Ex-Fin EBIT Margin", "Total EBIT Margin (coverage varies)"]
    for name in required_summary:
        if name not in sh:
            failures.append(f"Missing Growth Summary column: {name}")
    if not failures:
        for row, year in enumerate(range(2024, 2031), 5):
            if summary.cell(row, sh["Year"]).value != year:
                failures.append(f"Summary year mismatch at row {row}")
            expected = audit["summary"][str(year)]
            expected_ex = audit["summary_ex_financials"][str(year)]
            if not close(summary.cell(row, sh["Revenue USD bn"]).value, expected["revenue_usd_bn"]):
                failures.append(f"Revenue mismatch FY{year}")
            if not close(summary.cell(row, sh["Headline Ex-Fin EBIT Margin"]).value, expected_ex["weighted_ebit_margin"]):
                failures.append(f"Ex-financials EBIT margin mismatch FY{year}")
            if not close(summary.cell(row, sh["Total EBIT Margin (coverage varies)"]).value, expected["weighted_ebit_margin"]):
                failures.append(f"Total EBIT margin mismatch FY{year}")
            low_formula = summary.cell(row, sh["AI @ 2% USD bn"]).value
            high_formula = summary.cell(row, sh["AI @ 3% USD bn"]).value
            if "'Inputs & Anchors'!$B$5" not in str(low_formula) or "'Inputs & Anchors'!$B$6" not in str(high_formula):
                failures.append(f"AI formula reference mismatch FY{year}")

    company = wb["Company Model"]
    company_count = company.max_row - 4
    checks["company_rows"] = company_count
    if company_count != 4946:
        failures.append(f"Company rows {company_count} != 4946")
    ch = hmap(company)
    for label in ("AI_2030_2pct_USD_m", "AI_2030_3pct_USD_m"):
        col = ch.get(label)
        count = 0 if col is None else sum(
            isinstance(company.cell(r, col).value, str) and company.cell(r, col).value.startswith("=")
            for r in range(5, company.max_row + 1)
        )
        checks[f"{label}_formula_count"] = count
        if count != 4946:
            failures.append(f"{label} formula count {count} != 4946")

    formula_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("=") and "#REF!" in value:
                    formula_errors.append(f"{ws.title}!{cell.coordinate}")
    checks["ref_formula_errors"] = formula_errors
    if formula_errors:
        failures.append(f"#REF formulas: {formula_errors[:10]}")

    checks["by_index_rows"] = wb["By Index"].max_row - 4
    checks["by_sector_rows"] = wb["By Sector"].max_row - 4
    if checks["by_index_rows"] != 49:
        failures.append("By Index should have 7 universes x 7 years = 49 rows")
    if checks["by_sector_rows"] != len(audit["terminal_growth"]) * 7:
        failures.append("By Sector row count mismatch")

    sensitivity = audit["terminal_sensitivity"]["2030"]
    for excel_row, key in zip((16, 17, 18), ("normalized", "base", "bull")):
        if not close(summary.cell(excel_row, 2).value, sensitivity[key]["revenue_usd_bn"]):
            failures.append(f"2030 sensitivity revenue mismatch: {key}")
        if not close(summary.cell(excel_row, 5).value, sensitivity[key]["ex_financials_ebit_margin"]):
            failures.append(f"2030 sensitivity EBIT margin mismatch: {key}")

    bridge = audit["ebit_margin_bridge_2028"]
    if not close(summary.cell(23, 2).value, bridge["ex_financials"]["direct_revenue_coverage_of_modeled"]):
        failures.append("2028 ex-fin EBIT bridge coverage mismatch")
    if not close(summary.cell(24, 2).value, bridge["total"]["direct_revenue_coverage_of_modeled"]):
        failures.append("2028 total EBIT bridge coverage mismatch")

    it = audit["terminal_growth"]["Information Technology"]
    checks["it_growth_path_2028_2030"] = [it["consensus_2028"], it["base_2029"], it["base_2030"]]
    if it["base_2029"] > it["consensus_2028"] or it["base_2030"] > it["base_2029"]:
        failures.append("Information Technology growth reaccelerates in the neutral terminal path")
    if it["base_2030"] > 0.20:
        warnings.append("Information Technology terminal growth exceeds 20% and is load-bearing")

    s2030 = audit["summary"]["2030"]
    cagr = (s2030["revenue_usd_bn"] / audit["summary"]["2026"]["revenue_usd_bn"]) ** 0.25 - 1
    checks.update({
        "revenue_2026_2030_cagr": cagr,
        "revenue_2030_usd_bn": s2030["revenue_usd_bn"],
        "ai_2030_low_usd_bn": s2030["ai_low_usd_bn"],
        "ai_2030_high_usd_bn": s2030["ai_high_usd_bn"],
        "ex_fin_ebit_margin_2030": audit["summary_ex_financials"]["2030"]["weighted_ebit_margin"],
        "total_ebit_margin_2030": s2030["weighted_ebit_margin"],
        "skeptic_verdict": audit.get("skeptic_verdict"),
    })
    if cagr > 0.10:
        warnings.append("2026-2030 CAGR above 10%; inspect terminal assumptions")
    if sensitivity["base"]["revenue_usd_bn"] / sensitivity["normalized"]["revenue_usd_bn"] - 1 < 0.02:
        warnings.append("Historical-normalized case is close to base and should not be presented as a balanced downside")

    report = {"status": "PASS" if not failures else "FAIL", "failures": failures, "warnings": warnings, "checks": checks}
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
